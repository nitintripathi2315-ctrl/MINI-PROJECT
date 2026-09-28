from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

# --- Path-based versions (used by the original CLI script) ---


def read_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_path: Path) -> str:
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path: Path) -> str | None:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix == ".docx":
        return read_docx(file_path)
    else:
        return None


# --- Byte-based versions (used by the FastAPI upload endpoint) ---
# Same extraction logic, just reading from an in-memory buffer instead of
# a path on disk, since uploaded files never get saved to disk here.


def read_pdf_bytes(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx_bytes(content: bytes) -> str:
    document = Document(BytesIO(content))
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume_bytes(filename: str, content: bytes) -> str | None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return read_pdf_bytes(content)
    elif suffix == ".docx":
        return read_docx_bytes(content)
    else:
        return None
